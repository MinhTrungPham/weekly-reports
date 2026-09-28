import openpyxl
import os
import re

report = []

def check_excel_col(file_path, table_name, col_name, expected_type=None, expected_null=None, expected_key=None, expected_desc=None):
    if not os.path.exists(file_path):
        report.append(f"FAIL: {file_path} not found")
        return False
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    curr = ''
    found = False
    for row in range(1, ws.max_row + 1):
        c1 = str(ws.cell(row=row, column=1).value or '').strip()
        if c1: curr = c1
        c2 = str(ws.cell(row=row, column=2).value or '').strip()
        if curr == table_name and c2 == col_name:
            found = True
            c3 = str(ws.cell(row=row, column=3).value or '').strip()
            c4 = str(ws.cell(row=row, column=4).value or '').strip()
            c6 = str(ws.cell(row=row, column=6).value or '').strip()
            c7 = str(ws.cell(row=row, column=7).value or '').strip()
            
            if expected_type and expected_type not in c3:
                report.append(f"FAIL: {file_path} {table_name}.{col_name} Type '{c3}' != '{expected_type}'")
            if expected_null and expected_null not in c4:
                report.append(f"FAIL: {file_path} {table_name}.{col_name} Null '{c4}' != '{expected_null}'")
            if expected_key and expected_key not in c6:
                report.append(f"FAIL: {file_path} {table_name}.{col_name} Key '{c6}' != '{expected_key}'")
            if expected_desc and expected_desc not in c7:
                report.append(f"FAIL: {file_path} {table_name}.{col_name} Desc '{c7}' doesn't contain '{expected_desc}'")
            break
    if not found:
        report.append(f"FAIL: {file_path} {table_name}.{col_name} not found")
        return False
    return True

def check_puml(file_path, patterns):
    if not os.path.exists(file_path):
        report.append(f"FAIL: {file_path} not found")
        return
    with open(file_path, 'r', encoding='utf-8') as f:
        content = f.read()
    for p in patterns:
        if not re.search(p, content, re.MULTILINE):
            report.append(f"FAIL: {file_path} missing pattern: {p}")

# 1. user_roles
check_excel_col('Domain A/Data Dictionary domain A.xlsx', 'user_roles', 'id', 'BIGINT', 'N', 'PK')
check_excel_col('Domain A/Data Dictionary domain A.xlsx', 'user_roles', 'warehouse_id', None, 'Y', 'FK')
check_puml('Domain A/erd_domain_A.puml', [
    r'\*\s*id\s*:\s*BIGINT\s*<<PK>>',
    r'warehouse_id\s*:\s*BIGINT\s*<<FK nullable>>',
    r'UNIQUE\s*\(user_id,\s*role_id,\s*warehouse_id\)'
])

# 2. locations
check_excel_col('Domain B/Data_Dictionary_Domain_B.xlsx', 'locations', 'code', None, None, 'UK (with warehouse_id)')
check_puml('Domain B/ERD_Domain_B_Warehouse_Structure.puml', [
    r'\*\s*code\s*:\s*VARCHAR\(50\)\n', # NO UK
    r'UNIQUE\s*\(warehouse_id,\s*code\)',
    r'parent cung warehouse'
])

# 3. sku_uom_conversions
check_excel_col('Domain C/Data Dictionary domain C.xlsx', 'sku_uom_conversions', 'sku_id', None, None, 'UK (with uom_id)')
check_excel_col('Domain C/Data Dictionary domain C.xlsx', 'sku_uom_conversions', 'factor_to_base', None, None, None, 'Must be > 0')
check_puml('Domain C/erd_domain_C.puml', [
    r'UNIQUE\s*\(sku_id,\s*uom_id\)',
    r'CHECK\s*\(factor_to_base\s*>\s*0\)'
])

# 4. sku_barcodes
check_puml('Domain C/erd_domain_C.puml', [
    r'\*\s*barcode\s*:\s*VARCHAR\(100\)\s*<<UK>>',
    r'1 primary barcode cho moi \(sku_id,\s*uom_id\)'
])

# 5. sku_suppliers
check_excel_col('Domain C/Data Dictionary domain C.xlsx', 'sku_suppliers', 'is_preferred', 'BOOLEAN')
check_puml('Domain C/erd_domain_C.puml', [
    r'\*\s*is_preferred\s*:\s*BOOLEAN'
])

# 6. inventory
check_excel_col('Domain D/Data_Dictionary_Domain_D.xlsx', 'inventory', 'serial_no', 'VARCHAR(100)')
check_puml('Domain D/erd_domain_D.puml', [
    r'serial_no\s*:\s*VARCHAR\(100\)',
    r'UNIQUE\s*\(location_id,\s*sku_id,\s*lot_no,\s*serial_no\)',
    r'warehouse_id phai khop voi warehouse_id cua location_id'
])

# 7. stock_ledger
check_excel_col('Domain E/Data Dictionary domain E.xlsx', 'stock_ledger', 'serial_no', 'VARCHAR(100)')
check_puml('Domain E/erd_domain_E.puml', [
    r'serial_no\s*:\s*VARCHAR\(100\)',
    r'Reserve/Release khong ghi vao ledger',
    r'Ledger va inventory phai update trong cung transaction'
])

# 8. PICK / SHIP
check_puml('Domain E/erd_domain_E.puml', [
    r'PICK_OUT',
    r'PICK_IN',
    r'chi tru hang o staging'
])

# 9. Audit (check erd_all.puml)
check_puml('02_ERD/erd_all.puml', [
    r'users\s*\|\|--o\{\s*stock_ledger\s*:\s*"created_by"',
    r'users\s*\|\|--o\{\s*inventory\s*:\s*"created_by"'
])

if report:
    print('\n'.join(report))
else:
    print('ALL CHECKS PASSED')
